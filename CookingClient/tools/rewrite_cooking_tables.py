# -*- coding: utf-8 -*-
"""Rewrite CookingTables.xlsx for 七日食谱 DAY1-6 (+ DAY7 空火锅占位)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pathlib import Path

wb = Workbook()
HEADER_FILL = PatternFill("solid", fgColor="D9EAD3")


def sheet(name, headers, rows):
    if name in wb.sheetnames:
        ws = wb[name]
    elif wb.active.title == "Sheet":
        ws = wb.active
        ws.title = name
    else:
        ws = wb.create_sheet(name)

    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for row in rows:
        ws.append(list(row))

    for col in ws.columns:
        width = 12
        for cell in col:
            if cell.value is not None:
                width = max(width, min(52, len(str(cell.value)) + 2))
        ws.column_dimensions[col[0].column_letter].width = width
    return ws


readme = wb.active
readme.title = "说明"
readme["A1"] = "七日食谱表（CookingTables）"
readme["A1"].font = Font(bold=True, size=14)
for i, line in enumerate(
    [
        "",
        "导入：Cooking / Import Cooking Tables",
        "第一行必须是英文表头。",
        "shape：Whole/Dice/Slice/Shred；cookMethod：Raw/Boiled/Fried/BoiledAndFried",
        "RecipeSeasonings 行序 = 该食材上调料顺序。",
        "RecipeSteps 为日记文案；暂可不参与判定。",
        "DAY7 火锅：ingredients/seasonings 留空。",
    ],
    start=2,
):
    readme[f"A{i}"] = line
readme.column_dimensions["A"].width = 72

# ---------- Seasonings ----------
sheet(
    "Seasonings",
    ["id", "displayName", "icon"],
    [
        ("season_oil", "油", ""),
        ("season_salt", "盐", ""),
        ("season_pepper", "黑胡椒", ""),
        ("season_sugar", "糖", ""),
        ("season_sauce", "酱料", ""),
        ("season_spice", "香料", ""),
    ],
)

# ---------- Recipes ----------
sheet(
    "Recipes",
    ["id", "displayName", "unlockDay", "finishedSprite"],
    [
        ("recipe_day1", "什锦蘑菇浓汤", 1, ""),
        ("recipe_day2", "虹彩北非蛋", 2, ""),
        ("recipe_day3", "蜜莓酱香煎玉米鱼", 3, ""),
        ("recipe_day4", "幽灵布丁船", 4, ""),
        ("recipe_day5", "海鲜杂烩饭", 5, ""),
        ("recipe_day6", "鸟巢山药饼", 6, ""),
        ("recipe_day7", "七日火锅", 7, ""),
    ],
)

# ---------- RecipeIngredients ----------
# 依据步骤推断 shape / cookMethod
sheet(
    "RecipeIngredients",
    ["recipeId", "ingredientId", "shape", "cookMethod", "count"],
    [
        # DAY1：白蘑菇切片 → 炒至焦香 → 炖煮
        ("recipe_day1", "ing_mushroom", "Slice", "BoiledAndFried", 1),
        # DAY2：牛肉/洋葱/番茄切丁炒；鸡蛋焖煮
        ("recipe_day2", "ing_beef", "Dice", "Fried", 1),
        ("recipe_day2", "ing_onion", "Dice", "Fried", 1),
        ("recipe_day2", "ing_tomato", "Dice", "Fried", 1),
        ("recipe_day2", "ing_egg", "Whole", "Boiled", 1),
        # DAY3：玉米鱼切块香煎；蜜莓果切丁熬酱
        ("recipe_day3", "ing_cornfish", "Dice", "Fried", 1),
        ("recipe_day3", "ing_berry", "Dice", "Boiled", 1),
        # DAY4：南瓜切块煮泥；蜜莓装饰
        ("recipe_day4", "ing_pumpkin", "Dice", "Boiled", 1),
        ("recipe_day4", "ing_berry", "Whole", "Raw", 1),
        # DAY5：洋葱番茄炒酱；米炒；鱼虾焖
        ("recipe_day5", "ing_onion", "Dice", "Fried", 1),
        ("recipe_day5", "ing_tomato", "Dice", "Fried", 1),
        ("recipe_day5", "ing_rice", "Whole", "Fried", 1),
        ("recipe_day5", "ing_cornfish", "Dice", "Boiled", 1),
        ("recipe_day5", "ing_shrimp", "Whole", "Boiled", 1),
        # DAY6：山药丝煎饼；熟蛋丁/番茄丁/黄瓜片生拌
        ("recipe_day6", "ing_yam", "Shred", "Fried", 1),
        ("recipe_day6", "ing_doneegg", "Dice", "Raw", 1),
        ("recipe_day6", "ing_tomato", "Dice", "Raw", 1),
        ("recipe_day6", "ing_cucumber", "Slice", "Raw", 1),
        # DAY7：空
    ],
)

# ---------- RecipeSeasonings（按步骤落到具体食材；行序=顺序）----------
sheet(
    "RecipeSeasonings",
    ["recipeId", "seasoningId", "targetIngredientId", "count"],
    [
        # DAY1：锅中放油炒蘑菇 → 盐调味
        ("recipe_day1", "season_oil", "ing_mushroom", 1),
        ("recipe_day1", "season_salt", "ing_mushroom", 1),
        # DAY2：油炒洋葱/牛肉加盐；出锅前蛋上黑胡椒与香料
        ("recipe_day2", "season_oil", "ing_onion", 1),
        ("recipe_day2", "season_salt", "ing_onion", 1),
        ("recipe_day2", "season_oil", "ing_beef", 1),
        ("recipe_day2", "season_salt", "ing_beef", 1),
        ("recipe_day2", "season_pepper", "ing_egg", 1),
        ("recipe_day2", "season_spice", "ing_egg", 1),
        # DAY3：蜜莓加糖；鱼煎时油/盐/黑胡椒
        ("recipe_day3", "season_sugar", "ing_berry", 1),
        ("recipe_day3", "season_oil", "ing_cornfish", 1),
        ("recipe_day3", "season_salt", "ing_cornfish", 1),
        ("recipe_day3", "season_pepper", "ing_cornfish", 1),
        # DAY4：糖入南瓜泥；酱料画表情
        ("recipe_day4", "season_sugar", "ing_pumpkin", 1),
        ("recipe_day4", "season_sauce", "ing_pumpkin", 1),
        # DAY5：油炒洋葱番茄；米加盐；出锅前鱼虾黑胡椒
        ("recipe_day5", "season_oil", "ing_onion", 1),
        ("recipe_day5", "season_oil", "ing_tomato", 1),
        ("recipe_day5", "season_salt", "ing_rice", 1),
        ("recipe_day5", "season_pepper", "ing_cornfish", 1),
        ("recipe_day5", "season_pepper", "ing_shrimp", 1),
        # DAY6：山药丝盐胡椒后加油煎；沙拉酱+黑胡椒拌蛋蔬
        ("recipe_day6", "season_salt", "ing_yam", 1),
        ("recipe_day6", "season_pepper", "ing_yam", 1),
        ("recipe_day6", "season_oil", "ing_yam", 1),
        ("recipe_day6", "season_sauce", "ing_doneegg", 1),
        ("recipe_day6", "season_sauce", "ing_tomato", 1),
        ("recipe_day6", "season_sauce", "ing_cucumber", 1),
        ("recipe_day6", "season_pepper", "ing_doneegg", 1),
        ("recipe_day6", "season_pepper", "ing_tomato", 1),
        ("recipe_day6", "season_pepper", "ing_cucumber", 1),
    ],
)

# ---------- RecipeSteps（日记文案）----------
sheet(
    "RecipeSteps",
    ["recipeId", "order", "hintText", "action", "ingredientId", "cookMethod", "seasoningId"],
    [
        # DAY1
        ("recipe_day1", 1, "白蘑菇切片。", "Cut", "ing_mushroom", "Raw", ""),
        ("recipe_day1", 2, "锅中放油，加入蘑菇，炒至焦香。", "Fry", "ing_mushroom", "Fried", "season_oil"),
        ("recipe_day1", 3, "开始炖煮。", "Boil", "ing_mushroom", "BoiledAndFried", ""),
        ("recipe_day1", 4, "记得使用盐调味。", "Season", "ing_mushroom", "BoiledAndFried", "season_salt"),
        # DAY2
        ("recipe_day2", 1, "牛肉、洋葱、番茄切丁。", "Cut", "ing_beef", "Raw", ""),
        ("recipe_day2", 2, "锅中放油，分别加入洋葱和牛肉炒至出水，加盐。", "Fry", "ing_beef", "Fried", "season_oil"),
        ("recipe_day2", 3, "倒入番茄丁，收浓汤汁。", "Fry", "ing_tomato", "Fried", ""),
        ("recipe_day2", 4, "打入鸡蛋焖煮，出锅前加入黑胡椒与香料调味。", "Boil", "ing_egg", "Boiled", "season_pepper"),
        # DAY3
        ("recipe_day3", 1, "玉米鱼切块，蜜莓果切小丁。", "Cut", "ing_cornfish", "Raw", ""),
        ("recipe_day3", 2, "蜜莓果加糖，煮到浓稠。", "Boil", "ing_berry", "Boiled", "season_sugar"),
        ("recipe_day3", 3, "锅热放油，玉米鱼煎至两面金黄，加入盐和黑胡椒调味。", "Fry", "ing_cornfish", "Fried", "season_oil"),
        ("recipe_day3", 4, "鱼肉装盘，淋上蜜莓酱。", "Plate", "ing_cornfish", "Fried", ""),
        # DAY4
        ("recipe_day4", 1, "幽灵南瓜切块，煮熟冷却后加入奶油翻拌成泥。", "Boil", "ing_pumpkin", "Boiled", ""),
        ("recipe_day4", 2, "在盘子上摆出幽灵的形状，使用巧克力酱画出表情。", "Season", "ing_pumpkin", "Boiled", "season_sauce"),
        ("recipe_day4", 3, "在周围放置奶油、蜜莓果、糖果进行装饰。", "Plate", "ing_berry", "Raw", ""),
        # DAY5
        ("recipe_day5", 1, "鱼切块，处理好虾。", "Cut", "ing_cornfish", "Raw", ""),
        ("recipe_day5", 2, "锅中放油，依次加入洋葱与番茄，炒到出汁成酱。", "Fry", "ing_onion", "Fried", "season_oil"),
        ("recipe_day5", 3, "倒入大米翻炒，加入盐拌匀。", "Fry", "ing_rice", "Fried", "season_salt"),
        ("recipe_day5", 4, "铺上处理好的鱼和虾进行焖煮。", "Boil", "ing_cornfish", "Boiled", ""),
        ("recipe_day5", 5, "出锅前可再次进行调味。", "Season", "ing_shrimp", "Boiled", "season_pepper"),
        # DAY6
        ("recipe_day6", 1, "鸡蛋煮熟后切丁，番茄切丁、黄瓜切片，山药切丝。", "Cut", "ing_yam", "Raw", ""),
        ("recipe_day6", 2, "山药丝加入盐和黑胡椒调味，锅中加油，将山药丝煎成饼。", "Fry", "ing_yam", "Fried", "season_oil"),
        ("recipe_day6", 3, "将鸡蛋丁、番茄丁、黄瓜片与沙拉酱、黑胡椒拌匀。", "Season", "ing_doneegg", "Raw", "season_sauce"),
        ("recipe_day6", 4, "组装：取一张山药饼固定成碗状，中间放上拌好的沙拉。", "Plate", "ing_yam", "Fried", ""),
    ],
)

# ---------- Ingredients（保留原有）----------
ingredients = [
    ("ing_beef", "虹彩牛肉", "散发着彩虹光彩的牛肉，最美味的精华都浓缩在此", "Meat", False, True, 90, 160, "", "#FFFFFF"),
    ("ing_berry", "蜜莓果", "味道甘甜的果实，一口气吃太多小心牙疼。", "Vegetable", True, True, 70, 140, "", "#FFFFFF"),
    ("ing_cornfish", "玉米鱼", "通体黄色的有鳞鱼，肉质鲜甜，难怪被比作玉米。", "Meat", False, True, 110, 180, "", "#FFFFFF"),
    ("ing_cucumber", "黄瓜", "爽口的瓜类植物，为所有人带来清新香气的冲击。", "Vegetable", True, True, 60, 130, "", "#FFFFFF"),
    ("ing_doneegg", "熟鸡蛋", "白水煮熟的鸡蛋，蛋白包裹着饱满的蛋黄，完美！", "Meat", True, True, 100, 120, "", "#FFFFFF"),
    ("ing_egg", "鸡蛋", "不知道是哪种鸡下的蛋，看起来十分普通，希望味道也是。", "Meat", True, False, 100, 170, "", "#FFFFFF"),
    ("ing_mushroom", "白蘑菇", "生长在森林中的菌类，形似夜空中皎白的明月。", "Vegetable", False, True, 120, 180, "Assets/_Project/Arts/AI/Mushroom.png", "#FFFFFF"),
    ("ing_onion", "洋葱", "生长在森林中的野洋葱，味道辛甜。", "Vegetable", True, True, 110, 170, "", "#FFFFFF"),
    ("ing_potato", "星土豆", "外形和普通土豆差不多，切开后内部有星状的花纹。", "Vegetable", False, True, 130, 190, "", "#FFFFFF"),
    ("ing_pumpkin", "幽灵南瓜", "长有脸状花纹的白色南瓜，究竟是天生还是亡灵的执念呢？没有人知道答案。", "Vegetable", False, True, 125, 185, "", "#FFFFFF"),
    ("ing_rice", "大米", "一种形似米的谷物，味道也很接近，不如就当成米吧。", "Other", False, False, 140, 200, "", "#FFFFFF"),
    ("ing_shrimp", "红斑虾", "浑身长满斑纹的红褐色虾类，据说斑纹越多越美味。", "Meat", False, True, 95, 155, "", "#FFFFFF"),
    ("ing_tomato", "番茄", "艳红色的果实，微微的酸味让人胃口大开。", "Vegetable", True, True, 80, 150, "", "#FFFFFF"),
    ("ing_yam", "山药", "像长毛的树枝，有人说它是土豆的亲戚，反正我不信。", "Vegetable", False, True, 115, 175, "", "#FFFFFF"),
]
sheet(
    "Ingredients",
    ["id", "displayName", "description", "kind", "canEatRaw", "canCut", "optimalDoneness", "maxDoneness", "spriteWhole", "baseColor"],
    ingredients,
)

sheet(
    "IngredientVisualStates",
    ["ingredientId", "shape", "cookMethod", "sprite"],
    [],
)

order = [
    "说明",
    "Seasonings",
    "Recipes",
    "RecipeIngredients",
    "RecipeSeasonings",
    "RecipeSteps",
    "Ingredients",
    "IngredientVisualStates",
]
for i, name in enumerate(order):
    wb.move_sheet(name, offset=i - wb.sheetnames.index(name))

tables = Path(__file__).resolve().parents[1] / "Assets/_Project/Data/Tables"
out = tables / "CookingTables.xlsx"
try:
    wb.save(out)
except PermissionError:
    out = tables / "CookingTables_new.xlsx"
    wb.save(out)
    print("WARN: CookingTables.xlsx locked; wrote", out.name)

print("wrote", out)
print("recipes", wb["Recipes"].max_row - 1)
print("ingredients_req", wb["RecipeIngredients"].max_row - 1)
print("seasonings_req", wb["RecipeSeasonings"].max_row - 1)
print("steps", wb["RecipeSteps"].max_row - 1)
